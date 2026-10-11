#!/usr/bin/env python3
"""SommelAI-vragen: spreadsheet <-> assets/quiz-data.js

  python3 tools/vragen.py export   # quiz-data.js -> vragen/sommelai-vragen.xlsx
  python3 tools/vragen.py import   # vragen/sommelai-vragen.xlsx -> quiz-data.js

Spreadsheet: één tabblad per kleur (Rood, Wit). Per vraag één rij met de vraagtekst,
daaronder één rij per antwoordoptie met een x in de kolom van elke druif die een punt krijgt.
Kolom "type": basis (vaste eerste ronde), extra (wordt gekozen als hij helpt), duel (tweekeuze, telt dubbel).
Een nieuwe vraag toevoegen: kopieer het patroon, geef een nieuw id (kleine letters/cijfers/underscore).
"""
import json, re, sys, pathlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = pathlib.Path(__file__).resolve().parent.parent
JS = ROOT / 'assets' / 'quiz-data.js'
XLSX = ROOT / 'vragen' / 'sommelai-vragen.xlsx'
SHEETS = {'red': 'Rood', 'white': 'Wit'}

def load_js():
    s = JS.read_text(encoding='utf-8')
    return json.loads(re.search(r'window\.QUIZ_DATA = (.*);\s*$', s, re.S).group(1))

def save_js(d):
    head = JS.read_text(encoding='utf-8').split('window.QUIZ_DATA')[0]
    JS.write_text(head + 'window.QUIZ_DATA = ' + json.dumps(d, ensure_ascii=False, indent=1) + ';\n', encoding='utf-8')

def export():
    d = load_js(); wb = Workbook(); wb.remove(wb.active)
    for key, name in SHEETS.items():
        c = d[key]; ws = wb.create_sheet(name); grapes = c['grapes']
        ws.append(['id', 'type', 'vraag', 'antwoord'] + grapes + ['opmerking sommelier'])
        for cell in ws[1]: cell.font = Font(bold=True)
        fill = PatternFill('solid', fgColor='F0E9DC')
        core = set(c.get('core', []))
        rows = [(q, 'basis' if q['id'] in core else 'extra') for q in c['questions']] + [(q, 'duel') for q in c.get('duels', [])]
        for q, typ in rows:
            ws.append([q['id'], typ, q['q'], ''] + [''] * len(grapes))
            for cell in ws[ws.max_row]: cell.fill = fill; cell.font = Font(bold=True)
            for o in q['options']:
                ws.append(['', '', '', o['t']] + ['x' if g in o['g'] else '' for g in grapes])
        ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 8
        ws.column_dimensions['C'].width = 60; ws.column_dimensions['D'].width = 42
        for i in range(len(grapes)): ws.column_dimensions[chr(ord('E') + i)].width = 12
        ws.column_dimensions[chr(ord('E') + len(grapes))].width = 40
        ws.freeze_panes = 'E2'
        for row in ws.iter_rows(min_row=2):
            for cell in row[4:4 + len(grapes)]: cell.alignment = Alignment(horizontal='center')
    uitleg = wb.create_sheet('Uitleg', 0)
    for line in [
        'SommelAI-vragen — hoe werkt dit blad?',
        '',
        'Per kleur (Rood, Wit) staat elke vraag op een gekleurde rij, met daaronder de antwoordopties.',
        'Zet een x in de kolom van elke druif die een punt krijgt als iemand dat antwoord kiest. Mag meerdere druiven per antwoord zijn.',
        '',
        'type = basis: vaste eerste ronde, iedereen krijgt deze (8 vragen).',
        'type = extra: SommelAI stelt deze alleen als hij helpt om de koplopers uit elkaar te halen.',
        'type = duel: tweekeuzevraag die precies één druif tegenover één andere zet. Telt dubbel. Dit zijn de vragen die de twijfel oplossen.',
        '',
        'Je mag: teksten herschrijven, x-en verplaatsen, opties toevoegen of schrappen, hele vragen toevoegen (met nieuw id) of schrappen.',
        'Laat de kolom id staan voor bestaande vragen. Schrijf opmerkingen in de laatste kolom.',
        '',
        'De quiz: +1 per gekoppelde druif (duel +2). Hoogste score wint. Stopt zodra de koploper 3 punten voor staat, of na 20 vragen.',
    ]:
        uitleg.append([line])
    uitleg['A1'].font = Font(bold=True, size=14); uitleg.column_dimensions['A'].width = 120
    XLSX.parent.mkdir(exist_ok=True); wb.save(XLSX); print('geschreven', XLSX)

def import_():
    d = load_js(); wb = load_workbook(XLSX, data_only=True)
    for key, name in SHEETS.items():
        ws = wb[name]; rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
        grapes = [h for h in hdr[4:] if h and h != 'opmerking sommelier']
        questions, duels, core = [], [], []; cur = None
        for r in rows[1:]:
            qid, typ, qtext, opt = (r[0] or '').strip() if r[0] else '', (r[1] or '').strip().lower() if r[1] else '', r[2], r[3]
            if qid:
                cur = {'id': re.sub(r'[^a-z0-9_]', '', qid.lower()), 'q': str(qtext).strip(), 'options': []}
                if typ == 'duel': cur['duel'] = True; duels.append(cur)
                else:
                    questions.append(cur)
                    if typ == 'basis': core.append(cur['id'])
            elif cur and opt and str(opt).strip():
                cur['options'].append({'t': str(opt).strip(), 'g': [g for g, v in zip(grapes, r[4:4 + len(grapes)]) if v and str(v).strip().lower() == 'x']})
        for q in questions + duels:
            if not q['options']: sys.exit(f'Vraag {q["id"]} heeft geen antwoorden')
        d[key]['grapes'] = grapes; d[key]['questions'] = questions; d[key]['duels'] = duels; d[key]['core'] = core
        print(f'{name}: {len(questions)} vragen ({len(core)} basis), {len(duels)} duels')
    save_js(d); print('geschreven', JS)

if __name__ == '__main__':
    {'export': export, 'import': import_}[sys.argv[1]]()
